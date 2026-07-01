"""Format-specific parsers for non-PDF document ingestion.

Each loader returns ``PageResult`` objects: one page for single-unit formats,
or one per slide/sheet/chapter for multi-unit formats.
"""

import io
from pathlib import Path

from nepali_rag.ingestion.models import PageResult


def load_docx(file_path: str) -> list[PageResult]:
    """Extract paragraphs from a Word ``.docx`` file as a single page."""
    from docx import Document

    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)

    return [PageResult(
        page_number=1,
        text=text.strip(),
        source=str(Path(file_path).resolve()),
        extraction_method="docx"
    )]


def load_pptx(file_path: str) -> list[PageResult]:
    """Extract text from a PowerPoint ``.pptx`` file; one page per slide."""
    from pptx import Presentation

    prs = Presentation(file_path)
    results = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        texts.append(line)
        text = "\n".join(texts)
        results.append(PageResult(
            page_number=slide_num,
            text=text.strip(),
            source=str(Path(file_path).resolve()),
            extraction_method="pptx"
        ))

    return results

def load_rtf(file_path: str) -> list[PageResult]:
    """Extract plain text from an RTF file as a single page."""
    from striprtf.striprtf import rtf_to_text

    path = Path(file_path).resolve()
    raw = path.read_text(encoding="utf-8", errors="replace")
    text = rtf_to_text(raw)

    return [PageResult(
        page_number=1,
        text=text.strip(),
        source=str(path),
        extraction_method="rtf"
    )]


def load_html(file_path: str) -> list[PageResult]:
    """Extract visible text from an HTML file (strips scripts/styles) as a single page."""
    from bs4 import BeautifulSoup

    path = Path(file_path).resolve()
    raw = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")

    for tag in soup(["script", "style", "meta", "link"]):
        tag.decompose()

    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    text = "\n".join(lines)

    return [PageResult(
        page_number=1,
        text=text.strip(),
        source=str(path),
        extraction_method="html"
    )]


def load_epub(file_path: str) -> list[PageResult]:
    """Extract text from an EPUB ebook; one page per non-empty document item."""
    import ebooklib
    from ebooklib import epub
    from bs4 import BeautifulSoup

    book = epub.read_epub(file_path, options={"ignore_ncx": True})
    results = []
    page_num = 0

    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        soup = BeautifulSoup(item.get_content(), "html.parser")
        text = soup.get_text(separator="\n")
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        text = "\n".join(lines)

        if text.strip():
            page_num += 1
            results.append(PageResult(
                page_number=page_num,
                text=text.strip(),
                source=str(Path(file_path).resolve()),
                extraction_method="epub"
            ))

    return results


def load_odt(file_path: str) -> list[PageResult]:
    """Extract paragraphs from an OpenDocument ``.odt`` file as a single page."""
    from odf import text as odf_text
    from odf.opendocument import load as odf_load

    doc = odf_load(file_path)
    paragraphs = doc.getElementsByType(odf_text.P)
    texts = []
    for p in paragraphs:
        content = ""
        for node in p.childNodes:
            if hasattr(node, "data"):
                content += node.data
            elif hasattr(node, "__str__"):
                content += str(node)
        if content.strip():
            texts.append(content.strip())

    text = "\n\n".join(texts)

    return [PageResult(
        page_number=1,
        text=text.strip(),
        source=str(Path(file_path).resolve()),
        extraction_method="odt"
    )]
