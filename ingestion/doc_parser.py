import csv
import io
from pathlib import Path

from ingestion.models import PageResult


def load_docx(file_path: str) -> list[PageResult]:
    """Extract text from a .docx file. Each paragraph is concatenated into a single page."""
    from docx import Document

    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)

    return [PageResult(
        page_number=1,
        text=text.strip(),
        source=str(Path(file_path).resolve()),
        extraction_method="docx"
    )]


def load_pptx(file_path: str) -> list[PageResult]:
    """Extract text from a .pptx file. Each slide becomes a separate page."""
    from pptx import Presentation

    prs = Presentation(file_path)
    results = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        texts.append(line)
        text = "\n".join(texts)
        results.append(PageResult(
            page_number=slide_num,
            text=text.strip(),
            source=str(Path(file_path).resolve()),
            extraction_method="pptx"
        ))

    return results


def load_xlsx(file_path: str) -> list[PageResult]:
    """Extract text from a .xlsx file. Each sheet becomes a separate page."""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    results = []

    for sheet_num, sheet_name in enumerate(wb.sheetnames, start=1):
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            row_text = "\t".join(cells).strip()
            if row_text:
                rows.append(row_text)
        text = f"Sheet: {sheet_name}\n" + "\n".join(rows)
        results.append(PageResult(
            page_number=sheet_num,
            text=text.strip(),
            source=str(Path(file_path).resolve()),
            extraction_method="xlsx"
        ))

    wb.close()
    return results


def load_csv_file(file_path: str) -> list[PageResult]:
    """Extract text from a .csv file as a single page."""
    path = Path(file_path).resolve()

    with open(path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        rows = []
        for row in reader:
            rows.append("\t".join(row))
    text = "\n".join(rows)

    return [PageResult(
        page_number=1,
        text=text.strip(),
        source=str(path),
        extraction_method="csv"
    )]


def load_rtf(file_path: str) -> list[PageResult]:
    """Extract text from a .rtf file as a single page."""
    from striprtf.striprtf import rtf_to_text

    path = Path(file_path).resolve()
    raw = path.read_text(encoding="utf-8", errors="replace")
    text = rtf_to_text(raw)

    return [PageResult(
        page_number=1,
        text=text.strip(),
        source=str(path),
        extraction_method="rtf"
    )]


def load_html(file_path: str) -> list[PageResult]:
    """Extract text from an .html/.htm file as a single page."""
    from bs4 import BeautifulSoup

    path = Path(file_path).resolve()
    raw = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")

    for tag in soup(["script", "style", "meta", "link"]):
        tag.decompose()

    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    text = "\n".join(lines)

    return [PageResult(
        page_number=1,
        text=text.strip(),
        source=str(path),
        extraction_method="html"
    )]


def load_epub(file_path: str) -> list[PageResult]:
    """Extract text from an .epub file. Each chapter/item becomes a separate page."""
    import ebooklib
    from ebooklib import epub
    from bs4 import BeautifulSoup

    book = epub.read_epub(file_path, options={"ignore_ncx": True})
    results = []
    page_num = 0

    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        soup = BeautifulSoup(item.get_content(), "html.parser")
        text = soup.get_text(separator="\n")
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        text = "\n".join(lines)

        if text.strip():
            page_num += 1
            results.append(PageResult(
                page_number=page_num,
                text=text.strip(),
                source=str(Path(file_path).resolve()),
                extraction_method="epub"
            ))

    return results


def load_odt(file_path: str) -> list[PageResult]:
    """Extract text from an .odt file as a single page."""
    from odf import text as odf_text
    from odf.opendocument import load as odf_load

    doc = odf_load(file_path)
    paragraphs = doc.getElementsByType(odf_text.P)
    texts = []
    for p in paragraphs:
        content = ""
        for node in p.childNodes:
            if hasattr(node, "data"):
                content += node.data
            elif hasattr(node, "__str__"):
                content += str(node)
        if content.strip():
            texts.append(content.strip())

    text = "\n\n".join(texts)

    return [PageResult(
        page_number=1,
        text=text.strip(),
        source=str(Path(file_path).resolve()),
        extraction_method="odt"
    )]
